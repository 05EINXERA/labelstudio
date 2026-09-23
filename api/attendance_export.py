"""The attendance workbook and CSV.

After Q2 this is not a convenience feature. The two instances are never merged,
so manual comparison of exported files is the *only* cross-instance story —
which is why every sheet names the instance it came from, and why an
unidentified file would be useless.

**Durations are real numbers (decimal hours), never "7h 32m" strings.** This is
`api/routers/image_info.py`'s lesson applied: an attendance file exists to be
summed and pivoted, and "the recipient's first instinct is to select the column
and read the sum, and text cells make that silently produce nothing"
(`image_info.py:415-417`). A human-readable column may accompany the number; it
never replaces it.

Three more rules carried from that module, each for a reason it records:
dates as real dates so sorting and filtering work; a frozen header row and
auto-filter so the file is a working document rather than a printout; and an
explicit `number_format` so an empty cell reads `0.00` rather than a bare `0`
(openpyxl stores `float(0.0)` as an integer, so the cell *type* cannot keep a
column visually uniform — the format can).
"""
from datetime import datetime, timezone
from io import BytesIO, StringIO

import config
from api import attendance_report as report

# (label, width). Durations sit together so the numeric block is one
# contiguous range a reader can select and sum.
DAILY_COLUMNS = [
    ("Person", 20),
    ("Date", 12),
    ("First seen", 12),
    ("Latest seen", 12),
    ("End reason", 16),
    ("Sessions", 10),
    ("Present (h)", 13),
    ("Active timer (h)", 16),
    ("Breaks (h)", 12),
    ("Entered later (h)", 17),
    ("Tasks touched", 14),
    ("Tasks reviewed", 15),
    ("Instance", 18),
]

SESSION_COLUMNS = [
    ("Person", 20),
    ("Date", 12),
    ("Started", 12),
    ("Ended", 12),
    ("End reason", 16),
    ("Present (h)", 13),
    ("Breaks (h)", 12),
    ("Tasks touched", 14),
    ("Break detail", 42),
]

# Columns holding decimal hours, 1-based, per sheet. Used to apply the number
# format in one pass rather than cell by cell.
_DAILY_DURATION_COLS = (7, 8, 9, 10)
_SESSION_DURATION_COLS = (6, 7)

_END_REASON_LABEL = {
    "logout": "logged out",
    # Never "logged out at": a timeout end is the last moment the person was
    # seen, a lower bound. Rendering it as a logout time overstates the
    # precision of an attendance record.
    "timeout": "ended by timeout",
    "open": "still open",
}


def _hours(seconds) -> float:
    """Seconds as decimal hours, as a float.

    Rounded to 4 places rather than 2, for the reason `image_info.py` records:
    a short interval rounded to 2 places is 0.0, which openpyxl writes as the
    integer 0 — so the column becomes a mix of ints and floats and a genuinely
    short day reads as "no time at all". The cell's 0.00 display format still
    shows two places; the extra precision keeps the stored value honest and
    the SUM correct.
    """
    return float(round((seconds or 0) / 3600, 4))


def _local_time(moment) -> str:
    """A UTC instant as local clock time in the site zone.

    An attendance register is about when people were in the office, so the
    export shows office time wherever it is opened. Never a numeric offset:
    Kathmandu is +05:45.
    """
    if moment is None:
        return ""
    return report.as_utc(moment).astimezone(report.site_tz()).strftime("%H:%M")


def _end_reason(row) -> str:
    return _END_REASON_LABEL.get(row.get("last_seen_reason"), row.get("last_seen_reason") or "")


def _break_detail(session) -> str:
    """One session's breaks, flattened for the Sessions sheet.

    Text, deliberately: this column is read, not summed. The numeric break
    total is its own column beside it.
    """
    parts = []
    for brk in session.get("breaks", []):
        label = f"{_local_time(brk['started_at'])}-{_local_time(brk['ended_at'])}"
        if brk.get("source") == "manual":
            label += " (entered later)"
        if not brk.get("ended", True):
            label += " (not ended)"
        parts.append(label)
    return "; ".join(parts)


def build_workbook(rows, date_from, date_to, sessions_by_key=None) -> bytes:
    """The three-sheet workbook, as bytes.

    Imported inside the function rather than at module scope, against
    CLAUDE.md rule 2, for the reason `image_info.py:383-390` documents for the
    same dependency: openpyxl is pulled in by exactly one endpoint that is used
    occasionally, and the alternative is paying its import on every worker
    start for a feature most requests never touch. The rule's target is an
    import hidden inside a hot function to dodge a cycle; this is a deliberate,
    documented deferral of a heavy optional dependency.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="374151")

    def write_header(sheet, columns):
        sheet.append([label for label, _ in columns])
        for idx, (_, width) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[get_column_letter(idx)].width = width

    def finish(sheet, columns, duration_cols):
        for col in duration_cols:
            for row_cells in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                for cell in row_cells:
                    cell.number_format = "0.00"
        # A working document, not a printout.
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"
            )

    wb = Workbook()

    # --- Sheet 1: Summary ---------------------------------------------------
    # First, because it is the half that gets read and the half that gets
    # pasted into an email. It is also what identifies the file: without the
    # instance name, two exports from the two deployments are indistinguishable
    # and the only cross-instance mechanism there is stops working.
    summary = wb.active
    summary.title = "Summary"

    summary.append(["Attendance register"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Instance", config.ATTENDANCE_INSTANCE_ID])
    summary.append(["Timezone", config.ATTENDANCE_TZ])
    summary.append(["From", str(date_from)])
    summary.append(["To", str(date_to)])
    summary.append([
        "Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    ])
    summary.append([])

    people = {row["user_id"] for row in rows}
    days = {row["local_date"] for row in rows}
    summary.append(["People", len(people)])
    summary.append(["Days covered", len(days)])

    # The duration stats need the 0.00 format applied individually, because
    # they are a column of label/value pairs rather than a table with a header
    # row the bulk pass below can key off.
    #
    # The format is what keeps them uniform, NOT the cell type: openpyxl
    # narrows float(0.0) to an integer on write, so a zero total would render
    # as a bare "0" beside a neighbouring "0.25" without it. This is the exact
    # point image_info.py records for the same reason, and it is only visible
    # on a real day where someone's present time is zero.
    for label, seconds in (
        ("Present (h)", sum(r["present_seconds"] for r in rows)),
        ("Active timer (h)", sum(r["active_seconds"] for r in rows)),
        ("Breaks (h)", sum(r["break_seconds"] for r in rows)),
        ("Entered later (h)", sum(r["manual_break_seconds"] for r in rows)),
    ):
        summary.append([label, _hours(seconds)])
        summary.cell(row=summary.max_row, column=2).number_format = "0.00"

    summary.append(["Tasks touched", sum(r["tasks_touched"] for r in rows)])
    summary.append(["Tasks reviewed", sum(r["tasks_reviewed"] for r in rows)])
    summary.append([])

    # Per-day totals: what makes this a report rather than a data dump, and the
    # table that gets compared against the other instance's file.
    day_header_row = summary.max_row + 1
    summary.append(["Date", "People", "Present (h)", "Breaks (h)", "Tasks touched"])
    for cell in summary[day_header_row]:
        cell.font = header_font
        cell.fill = header_fill

    for day in sorted(days):
        day_rows = [r for r in rows if r["local_date"] == day]
        summary.append([
            day,
            len({r["user_id"] for r in day_rows}),
            _hours(sum(r["present_seconds"] for r in day_rows)),
            _hours(sum(r["break_seconds"] for r in day_rows)),
            sum(r["tasks_touched"] for r in day_rows),
        ])

    for col in (3, 4):
        for row_cells in summary.iter_rows(
            min_row=day_header_row + 1, min_col=col, max_col=col
        ):
            for cell in row_cells:
                cell.number_format = "0.00"

    summary.append([])
    # The qualifications belong in the file, not only in the UI. A register
    # that is misread is worse than one nobody reads, and this file will be
    # opened by people who never saw the dashboard's footnote.
    for note in (
        "Present time excludes declared breaks.",
        "Active timer is time the annotation timer was running; it is always "
        "lower than Present and the two will not agree.",
        "An 'ended by timeout' day was not logged out of. The time shown is "
        "the last moment the person was seen — a lower bound, not a departure "
        "time.",
        "'Entered later' is break time added from the profile page after the "
        "fact rather than declared at the time. It is provenance, not a flag "
        "of doubt.",
        "Tasks touched is tasks the person had open. It is NOT tasks "
        "completed, which this system cannot attribute to a person.",
    ):
        summary.append([note])

    for label, width in (("A", 26), ("B", 16), ("C", 14), ("D", 14), ("E", 16)):
        summary.column_dimensions[label].width = width

    # --- Sheet 2: Daily -----------------------------------------------------
    daily = wb.create_sheet("Daily")
    write_header(daily, DAILY_COLUMNS)
    for row in rows:
        daily.append([
            row["username"],
            row["local_date"],          # a real date, so sorting works
            _local_time(row["first_seen"]),
            _local_time(row["last_seen"]),
            _end_reason(row),
            row["session_count"],
            _hours(row["present_seconds"]),
            _hours(row["active_seconds"]),
            _hours(row["break_seconds"]),
            _hours(row["manual_break_seconds"]),
            row["tasks_touched"],
            row["tasks_reviewed"],
            config.ATTENDANCE_INSTANCE_ID,
        ])
    for row_cells in daily.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row_cells:
            cell.number_format = "yyyy-mm-dd"
    finish(daily, DAILY_COLUMNS, _DAILY_DURATION_COLS)

    # --- Sheet 3: Sessions --------------------------------------------------
    # The sheet that answers "why does this day look odd" — the popup's data,
    # flat.
    sessions_sheet = wb.create_sheet("Sessions")
    write_header(sessions_sheet, SESSION_COLUMNS)
    for row in rows:
        key = (row["user_id"], row["local_date"])
        for session in (sessions_by_key or {}).get(key, []):
            sessions_sheet.append([
                row["username"],
                row["local_date"],
                _local_time(session["started_at"]),
                _local_time(session["ended_at"]),
                _END_REASON_LABEL.get(
                    session["end_reason"], session["end_reason"]
                ),
                _hours(session["seconds"]),
                _hours(session["break_seconds"]),
                session["tasks_touched"],
                _break_detail(session),
            ])
    for row_cells in sessions_sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row_cells:
            cell.number_format = "yyyy-mm-dd"
    finish(sessions_sheet, SESSION_COLUMNS, _SESSION_DURATION_COLS)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(rows, date_from, date_to) -> str:
    """The Daily sheet, flat.

    A bare download rather than a ZIP, matching how `exports.py` already treats
    CSV. The instance and timezone ride in comment lines above the header: a
    CSV cannot carry a Summary sheet, and an export that does not say which
    deployment produced it is useless for the one job these files have.
    """
    import csv

    out = StringIO()
    out.write(f"# Attendance register\n")
    out.write(f"# Instance,{config.ATTENDANCE_INSTANCE_ID}\n")
    out.write(f"# Timezone,{config.ATTENDANCE_TZ}\n")
    out.write(f"# From,{date_from},To,{date_to}\n")
    out.write(
        f"# Generated (UTC),{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}\n"
    )
    out.write("# Present excludes declared breaks. 'ended by timeout' is a lower bound, not a departure time.\n")
    out.write("# Tasks touched is NOT tasks completed, which cannot be attributed to a person.\n")

    writer = csv.writer(out, lineterminator="\n")
    writer.writerow([label for label, _ in DAILY_COLUMNS])
    for row in rows:
        writer.writerow([
            row["username"],
            row["local_date"],
            _local_time(row["first_seen"]),
            _local_time(row["last_seen"]),
            _end_reason(row),
            row["session_count"],
            # Decimal hours here too, for the same reason: a CSV is opened in a
            # spreadsheet and summed.
            _hours(row["present_seconds"]),
            _hours(row["active_seconds"]),
            _hours(row["break_seconds"]),
            _hours(row["manual_break_seconds"]),
            row["tasks_touched"],
            row["tasks_reviewed"],
            config.ATTENDANCE_INSTANCE_ID,
        ])
    return out.getvalue()
