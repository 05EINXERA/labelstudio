"""Spreadsheet rendering for the Image Inventory download.

Separate from `image_inventory` (which owns the queries) so the workbook
formatting is testable on its own, and so the heavy openpyxl import lands in
one place.

Built synchronously rather than through the export job queue: there is no
image rasterisation here, just text, and adding it to the in-process job
registry would grow shared state for no benefit (that registry cannot be
sharded across workers -- see CLAUDE.md rule 9).
"""
import logging
import re
from datetime import datetime, timezone
from io import BytesIO
from typing import List, Optional

from schemas import IMAGE_SIZE_CATEGORIES

logger = logging.getLogger(__name__)

# Characters Excel rejects in a sheet name, plus its length cap. A project
# called "Site A / Roofs [2026]" otherwise produces a file Excel refuses to
# open -- which users reasonably report as "the export is broken" rather than
# as a naming problem.
_SHEET_NAME_STRIP = re.compile(r"[\[\]:*?/\\]+")
_SHEET_NAME_MAX = 31

# Filesystem/header-hostile characters in a download filename. Mirrors
# team._safe_filename_part; the value goes into a quoted Content-Disposition
# header.
_FILENAME_STRIP = re.compile(r"[^A-Za-z0-9._-]+")


def safe_sheet_name(name: str, fallback: str = "Images") -> str:
    """An Excel-legal worksheet name derived from a project name."""
    cleaned = _SHEET_NAME_STRIP.sub("-", name or "").strip()
    # Excel also rejects a name that starts or ends with an apostrophe.
    cleaned = cleaned.strip("'").strip()
    cleaned = cleaned[:_SHEET_NAME_MAX]
    return cleaned or fallback


def safe_filename_part(value: str, fallback: str = "project") -> str:
    """Reduce a project name to characters safe in a Content-Disposition header."""
    cleaned = _FILENAME_STRIP.sub("_", value or "").strip("_")
    return cleaned[:60] or fallback


def download_filename(project_name: str) -> str:
    """A filename carrying the project and the date."""
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    return f"{safe_filename_part(project_name)}-images-{stamp}.xlsx"


def build_workbook(project_name: str, rows: List[dict], summary: dict) -> bytes:
    """Render the inventory as a real .xlsx file.

    `rows` are dicts with the same keys as ImageInventoryRow.
    """
    # Imported here rather than at module top: openpyxl is heavy (it pulls in
    # its own XML machinery) and is used by exactly this one endpoint, which is
    # rarely called. Deferring it keeps it out of every worker's import path
    # and off the startup cost of a deployment that may never download a
    # spreadsheet. This is a deliberate exception to the imports-at-top rule,
    # noted here so the next reader does not "fix" it.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_sheet_name(project_name)

    headers = ["Filename", "Width", "Height", "Resolution", "Size category",
               "File size (bytes)", "Status"]
    sheet.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F8B8D")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        width = row.get("width")
        height = row.get("height")
        resolution = f"{width} x {height}" if width and height else ""
        sheet.append([
            row.get("filename") or "",
            width,
            height,
            resolution,
            row.get("category") or "",
            row.get("file_size"),
            row.get("status") or "",
        ])

    # File size as a real number, not a formatted string. The recipient's first
    # instinct is to select the column and read the sum, and a text column
    # silently produces nothing. A missing file stays empty rather than
    # becoming 0, so the dash in the web view and the blank here agree.
    size_column = headers.index("File size (bytes)") + 1
    for excel_row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=excel_row, column=size_column)
        if cell.value is None:
            continue
        # Written as a number and never as a preformatted string; the number
        # format below handles presentation.
        #
        # Byte counts are whole numbers, so they are written as int rather than
        # float. Do not "normalise" this to float() thinking it guarantees a
        # float cell -- openpyxl serialises an integral float straight back to
        # an integer (float(4096) round-trips as 4096), so the cast achieves
        # nothing while suggesting it does. Both are the same numeric type in
        # xlsx and SUM() reads either; the only value that breaks the sum is a
        # string. The related trap the spec warns about -- rounding to two
        # places yielding a mixed int/float column -- cannot arise here because
        # no rounding happens: these are exact byte counts from os.path.getsize.
        cell.value = int(cell.value)
        cell.number_format = "#,##0"

    widths = {"Filename": 46, "Width": 10, "Height": 10, "Resolution": 16,
              "Size category": 15, "File size (bytes)": 18, "Status": 14}
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[header]

    # Freeze the header and turn on filter dropdowns over the data range.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"

    _add_summary_sheet(workbook, summary, safe_sheet_name(project_name))

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _add_summary_sheet(workbook, summary: dict, data_sheet_name: str) -> None:
    """Per-category counts on a second sheet -- the half that gets emailed."""
    from openpyxl.styles import Font, PatternFill

    name = safe_sheet_name("Summary")
    if name == data_sheet_name:
        name = safe_sheet_name("Summary (counts)")
    sheet = workbook.create_sheet(name)

    sheet.append(["Size category", "Images"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F8B8D")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    counts = {c["category"]: c["count"] for c in summary.get("categories", [])}
    # Every category, including zeroes, matching the web view.
    for category in IMAGE_SIZE_CATEGORIES:
        sheet.append([category, counts.get(category, 0)])

    sheet.append([])
    sheet.append(["Total images", summary.get("total", 0)])

    total_size = summary.get("total_size")
    label = "Total size (bytes)"
    if summary.get("total_size_is_complete") is False:
        # Never present a partial figure as a project total.
        label = "Total size (bytes, listed rows only)"
    sheet.append([label, int(total_size) if total_size is not None else None])

    sheet.append(["Files missing from disk", summary.get("missing_files", 0)])

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 18
